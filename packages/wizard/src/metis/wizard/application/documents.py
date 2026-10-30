"""Read bounded text from request-scoped Interview document uploads."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, BinaryIO

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader
from pypdf.generic import ArrayObject, DictionaryObject, StreamObject

from metis.wizard.contracts.documents import (
    DocumentExtractionErrorCode,
    DocumentExtractionResult,
    ExtractedDocument,
)
from metis.wizard.domain.documents import (
    EXTRACTION_DOCUMENT_EXTENSIONS,
    MAX_DOCUMENT_TEXT_CHARACTERS,
    MAX_DOCUMENTS,
    MAX_DOCX_ARCHIVE_ENTRIES,
    MAX_DOCX_EXPANDED_BYTES,
    MAX_EXTRACTION_FILE_BYTES,
    MAX_EXTRACTION_REQUEST_BYTES,
    MAX_PDF_EXPANDED_BYTES,
    normalize_fixture_name,
)

if TYPE_CHECKING:
    from collections.abc import Iterable

    from pypdf.generic import PdfObject

    from metis.wizard.application.compile.specification import SpecificationFixture

_READ_CHUNK_SIZE = 64 * 1024


class DocumentExtractionError(ValueError):
    """A user-correctable extraction failure with a stable code and file name."""

    def __init__(
        self,
        code: DocumentExtractionErrorCode,
        message: str,
        *,
        filename: str = "document",
    ) -> None:
        super().__init__(message)
        self.code = code
        self.filename = filename


class _TextLimitReached(Exception):
    """Stop a parser as soon as it emits more text than the contract permits."""


def extract_documents(
    uploads: Iterable[SpecificationFixture],
) -> DocumentExtractionResult:
    """Extract all files or raise one typed error; never return partial results."""
    items = list(uploads)
    if not items:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.EMPTY_DOCUMENT,
            "Select at least one document to read.",
        )
    if len(items) > MAX_DOCUMENTS:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.TOO_MANY_DOCUMENTS,
            f"Read at most {MAX_DOCUMENTS} documents at a time.",
        )

    seen: set[str] = set()
    streams: list[tuple[str, str, BinaryIO]] = []
    for item in items:
        try:
            name = normalize_fixture_name(item.name)
        except ValueError as exc:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.INVALID_FILENAME,
                str(exc),
                filename=item.name,
            ) from exc
        extension = Path(name).suffix.lower().removeprefix(".")
        if extension not in EXTRACTION_DOCUMENT_EXTENSIONS:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.UNSUPPORTED_FILE_TYPE,
                "Interview text is supported only for TXT, Markdown, YAML, JSON, CSV, PDF, and DOCX files.",
                filename=name,
            )
        identity = name.casefold()
        if identity in seen:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.INVALID_FILENAME,
                "The upload contains the same file name more than once.",
                filename=name,
            )
        seen.add(identity)
        streams.append((name, extension, item.stream))

    output: list[ExtractedDocument] = []
    total_bytes = 0
    for name, extension, stream in streams:
        data, total_bytes = _read_upload(stream, total_bytes, name)
        if extension in {"txt", "md", "yaml", "yml", "json", "csv"}:
            text = _extract_text_file(data, name)
        elif extension == "pdf":
            text = _extract_pdf(data, name)
        else:
            text = _extract_docx(data, name)
        if len(text) > MAX_DOCUMENT_TEXT_CHARACTERS:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.TEXT_TOO_LARGE,
                f"Extracted text exceeds {MAX_DOCUMENT_TEXT_CHARACTERS:,} characters.",
                filename=name,
            )
        if not text.strip():
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.EMPTY_DOCUMENT,
                "No readable text was found. Scanned PDFs need OCR, which is not supported.",
                filename=name,
            )
        output.append(ExtractedDocument(name=name, text=text))

    return DocumentExtractionResult(documents=output)


def _read_upload(stream: BinaryIO, request_bytes: int, name: str) -> tuple[bytes, int]:
    content = bytearray()
    while chunk := stream.read(_READ_CHUNK_SIZE):
        content.extend(chunk)
        if len(content) > MAX_EXTRACTION_FILE_BYTES:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.FILE_TOO_LARGE,
                f"Each document must be at most {MAX_EXTRACTION_FILE_BYTES:,} bytes.",
                filename=name,
            )
        if request_bytes + len(content) > MAX_EXTRACTION_REQUEST_BYTES:
            raise DocumentExtractionError(
                DocumentExtractionErrorCode.REQUEST_TOO_LARGE,
                f"Documents must be at most {MAX_EXTRACTION_REQUEST_BYTES:,} bytes in total.",
                filename=name,
            )
    if not content:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.EMPTY_DOCUMENT,
            "The file is empty.",
            filename=name,
        )
    return bytes(content), request_bytes + len(content)


def _extract_text_file(data: bytes, name: str) -> str:
    try:
        return data.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.MALFORMED_FILE,
            "Text files must use UTF-8 encoding.",
            filename=name,
        ) from exc


def _extract_pdf(data: bytes, name: str) -> str:
    try:
        reader = PdfReader(
            io.BytesIO(data), strict=False, root_object_recovery_limit=10_000
        )
        try:
            if reader.is_encrypted:
                raise DocumentExtractionError(
                    DocumentExtractionErrorCode.ENCRYPTED_PDF,
                    "Encrypted PDFs are not supported.",
                    filename=name,
                )
            parts: list[str] = []
            total_characters = 0
            # Check every reachable stream, including nested forms and font resources,
            # before text extraction can expand them outside the page content stream.
            pending: list[PdfObject] = [reader.trailer]
            seen: set[int] = set()
            total_expanded = 0
            while pending:
                item = pending.pop().get_object()
                if id(item) in seen:
                    continue
                seen.add(id(item))
                if isinstance(item, StreamObject):
                    total_expanded += len(item.get_data())
                    if total_expanded > MAX_PDF_EXPANDED_BYTES:
                        raise DocumentExtractionError(
                            DocumentExtractionErrorCode.FILE_TOO_LARGE,
                            f"PDF expanded content must be at most {MAX_PDF_EXPANDED_BYTES:,} bytes.",
                            filename=name,
                        )
                if isinstance(item, DictionaryObject):
                    pending.extend(item.values())
                elif isinstance(item, ArrayObject):
                    pending.extend(item)

            def collect(fragment: str, *_: object) -> None:
                nonlocal total_characters
                total_characters += len(fragment)
                if total_characters > MAX_DOCUMENT_TEXT_CHARACTERS:
                    raise _TextLimitReached
                parts.append(fragment)

            for page in reader.pages:
                page.extract_text(visitor_text=collect)
            return "".join(parts)
        finally:
            reader.close()
    except DocumentExtractionError:
        raise
    except _TextLimitReached as exc:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.TEXT_TOO_LARGE,
            f"Extracted text exceeds {MAX_DOCUMENT_TEXT_CHARACTERS:,} characters.",
            filename=name,
        ) from exc
    except Exception as exc:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.MALFORMED_FILE,
            "The PDF could not be read.",
            filename=name,
        ) from exc


def _extract_docx(data: bytes, name: str) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ARCHIVE_ENTRIES:
                raise DocumentExtractionError(
                    DocumentExtractionErrorCode.FILE_TOO_LARGE,
                    f"DOCX files can contain at most {MAX_DOCX_ARCHIVE_ENTRIES:,} archive entries.",
                    filename=name,
                )
            expanded = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise DocumentExtractionError(
                        DocumentExtractionErrorCode.MALFORMED_FILE,
                        "Encrypted DOCX files are not supported.",
                        filename=name,
                    )
                expanded += _check_archive_entry(archive, entry, expanded, name)
        document = Document(io.BytesIO(data))
        parts: list[str] = []
        total_characters = 0

        def add(text: str) -> None:
            nonlocal total_characters
            total_characters += len(text) + 1
            if total_characters > MAX_DOCUMENT_TEXT_CHARACTERS:
                raise _TextLimitReached
            parts.append(text)

        def add_table(table: Table) -> None:
            for row in table.rows:
                for cell in row.cells:
                    for block in cell.iter_inner_content():
                        if isinstance(block, Paragraph):
                            add(block.text)
                        elif isinstance(block, Table):
                            add_table(block)

        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                add(block.text)
            elif isinstance(block, Table):
                add_table(block)
        return "\n".join(parts)
    except DocumentExtractionError:
        raise
    except _TextLimitReached as exc:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.TEXT_TOO_LARGE,
            f"Extracted text exceeds {MAX_DOCUMENT_TEXT_CHARACTERS:,} characters.",
            filename=name,
        ) from exc
    except Exception as exc:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.MALFORMED_FILE,
            "The DOCX file could not be read.",
            filename=name,
        ) from exc


def _check_archive_entry(
    archive: zipfile.ZipFile,
    entry: zipfile.ZipInfo,
    expanded: int,
    name: str,
) -> int:
    if expanded + entry.file_size > MAX_DOCX_EXPANDED_BYTES:
        raise DocumentExtractionError(
            DocumentExtractionErrorCode.FILE_TOO_LARGE,
            f"DOCX expanded content must be at most {MAX_DOCX_EXPANDED_BYTES:,} bytes.",
            filename=name,
        )
    measured = 0
    with archive.open(entry) as member:
        while chunk := member.read(_READ_CHUNK_SIZE):
            measured += len(chunk)
            if expanded + measured > MAX_DOCX_EXPANDED_BYTES:
                raise DocumentExtractionError(
                    DocumentExtractionErrorCode.FILE_TOO_LARGE,
                    f"DOCX expanded content must be at most {MAX_DOCX_EXPANDED_BYTES:,} bytes.",
                    filename=name,
                )
    return measured


__all__ = ["DocumentExtractionError", "extract_documents"]
