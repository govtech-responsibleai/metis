"""The bounded text extraction contract for uploaded Interview documents."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)

from metis.wizard.application import documents
from metis.wizard.application.compile import SpecificationFixture
from metis.wizard.application.documents import (
    DocumentExtractionError,
    extract_documents,
)
from metis.wizard.contracts.documents import DocumentExtractionErrorCode


def _fixture(name: str, content: bytes) -> SpecificationFixture:
    return SpecificationFixture(name=name, stream=io.BytesIO(content))


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Policy opening")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Policy table"
    table.cell(0, 1).text = "Value"
    document.add_paragraph("Policy ending")
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def _pdf(text: str, *, form_depth: int = 0) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
    )
    page[NameObject("/Resources")] = writer._add_object(
        DictionaryObject(
            {
                NameObject("/Font"): writer._add_object(
                    DictionaryObject({NameObject("/F1"): font})
                )
            }
        )
    )
    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode())
    resources = page[NameObject("/Resources")]
    for _ in range(form_depth):
        content[NameObject("/Subtype")] = NameObject("/Form")
        content[NameObject("/Resources")] = resources
        form = writer._add_object(content.flate_encode())
        resources = DictionaryObject(
            {NameObject("/XObject"): DictionaryObject({NameObject("/Form1"): form})}
        )
        content = DecodedStreamObject()
        content.set_data(b"/Form1 Do")
    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = writer._add_object(content)
    result = io.BytesIO()
    writer.write(result)
    return result.getvalue()


def test_extracts_text_files_and_docx_body_in_source_order() -> None:
    result = extract_documents(
        [
            _fixture("notes.md", b"# Notes\nPolicy text"),
            _fixture("policy.docx", _docx()),
        ]
    )

    assert [document.name for document in result.documents] == [
        "notes.md",
        "policy.docx",
    ]
    assert result.documents[0].text == "# Notes\nPolicy text"
    assert (
        result.documents[1].text.index("Policy opening")
        < result.documents[1].text.index("Policy table")
        < result.documents[1].text.index("Policy ending")
    )


def test_extracts_pdf_text_and_rejects_a_pdf_without_text() -> None:
    result = extract_documents([_fixture("policy.pdf", _pdf("Policy text"))])
    assert "Policy text" in result.documents[0].text

    with pytest.raises(DocumentExtractionError) as caught:
        extract_documents([_fixture("scan.pdf", _pdf(""))])
    assert caught.value.code is DocumentExtractionErrorCode.EMPTY_DOCUMENT
    assert caught.value.filename == "scan.pdf"


@pytest.mark.parametrize("form_depth", [0, 1, 2])
def test_pdf_streams_share_an_expansion_limit(
    monkeypatch: pytest.MonkeyPatch, form_depth: int
) -> None:
    content = _pdf("Policy text" + " " * 2000, form_depth=form_depth)
    assert (
        "Policy text"
        in extract_documents([_fixture("policy.pdf", content)]).documents[0].text
    )
    monkeypatch.setattr(documents, "MAX_PDF_EXPANDED_BYTES", 1000)
    with pytest.raises(DocumentExtractionError) as caught:
        extract_documents([_fixture("policy.pdf", content)])
    assert caught.value.code is DocumentExtractionErrorCode.FILE_TOO_LARGE
    assert caught.value.filename == "policy.pdf"


def test_rejects_invalid_utf8_and_encrypted_pdfs() -> None:
    with pytest.raises(DocumentExtractionError) as invalid_text:
        extract_documents([_fixture("notes.txt", b"\xff")])
    assert invalid_text.value.code is DocumentExtractionErrorCode.MALFORMED_FILE

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    encrypted = io.BytesIO()
    writer.write(encrypted)
    with pytest.raises(DocumentExtractionError) as encrypted_pdf:
        extract_documents([_fixture("locked.pdf", encrypted.getvalue())])
    assert encrypted_pdf.value.code is DocumentExtractionErrorCode.ENCRYPTED_PDF


def test_rejects_malformed_and_oversized_documents(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    with pytest.raises(DocumentExtractionError) as malformed:
        extract_documents([_fixture("broken.docx", b"not a docx")])
    assert malformed.value.code is DocumentExtractionErrorCode.MALFORMED_FILE

    monkeypatch.setattr(documents, "MAX_EXTRACTION_FILE_BYTES", 4)
    with pytest.raises(DocumentExtractionError) as too_large:
        extract_documents([_fixture("large.txt", b"12345")])
    assert too_large.value.code is DocumentExtractionErrorCode.FILE_TOO_LARGE


def test_rejects_expanded_docx_archives_and_oversized_extracted_text(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    zip_bytes = io.BytesIO()
    with zipfile.ZipFile(zip_bytes, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 32)
    monkeypatch.setattr(documents, "MAX_DOCX_EXPANDED_BYTES", 16)
    with pytest.raises(DocumentExtractionError) as expanded:
        extract_documents([_fixture("large.docx", zip_bytes.getvalue())])
    assert expanded.value.code is DocumentExtractionErrorCode.FILE_TOO_LARGE

    monkeypatch.setattr(documents, "MAX_DOCUMENT_TEXT_CHARACTERS", 20)
    with pytest.raises(DocumentExtractionError) as text_limit:
        extract_documents([_fixture("large.txt", b"x" * 21)])
    assert text_limit.value.code is DocumentExtractionErrorCode.TEXT_TOO_LARGE


def test_extraction_rejects_a_batch_with_an_unreadable_file() -> None:
    with pytest.raises(DocumentExtractionError) as caught:
        extract_documents(
            [
                _fixture("first.txt", b"Readable"),
                _fixture("second.txt", b"\xff"),
            ]
        )

    assert caught.value.filename == "second.txt"
    assert caught.value.code is DocumentExtractionErrorCode.MALFORMED_FILE


def test_large_supporting_documents_keep_all_text() -> None:
    from metis.wizard import InterviewStartRequest
    from metis.wizard.contracts.interview_state import CapabilityConfidence

    texts = [("deep-dive.md", "a" * 307_000), ("support.yaml", "b" * 117_000)]
    extracted = extract_documents(
        [_fixture(name, text.encode()) for name, text in texts]
    )
    request = InterviewStartRequest(
        description="Evaluate the product",
        known_capability_ids=[],
        capability_confidence=CapabilityConfidence.LOW,
        documents=[item.model_dump() for item in extracted.documents],
    )
    assert [(item.name, item.text) for item in request.documents] == texts
